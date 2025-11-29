#FTO341
#Librerias
import openpyxl
from collections import Counter
from datetime import datetime
from pathlib import Path


def Generar_Plano_341():
    #Calcula el trimestre
    hoy = datetime.today()
    mes_actual = hoy.month
    año_actual = hoy.year
    
    if mes_actual in (1, 2, 3):
            trimestre_fin = datetime(año_actual - 1, 12, 31)
    elif mes_actual in (4, 5, 6):
                trimestre_fin = datetime(año_actual, 3, 31)
    elif mes_actual in (7, 8, 9):
                trimestre_fin = datetime(año_actual, 6, 30)
    else: trimestre_fin = datetime(año_actual, 9, 30)
    trimestre_str = trimestre_fin.strftime("%Y%m%d")    

   # BUSCA LA CARPETA
    carpeta = Path(fr"\\172.16.41.31\FIN-Contabilidad\04.FORMATOS REGULATORIOS\5._FORMATOS TRIMESTRALES\5._FTO 341\5._Formato 341 {trimestre_str}")
    if not carpeta.exists():
            print(f"No existe la carpeta del trimestre: {carpeta}")
            return
    
    # BUSCA EL ARCHIVO
    archivo_excel = next(carpeta.glob("F0000-104*.xlsm"), None)
    if not archivo_excel:
        print(f"No se encontró la proforma en la ruta {carpeta}")
        return
    
    #Routeros y config del archivo
    ruta_excel = archivo_excel   
    hoja_clientes = "Clientes"
    hoja_individual = "individual"
    fila_inicio = 6
    carpeta_salida = carpeta
    
    Path(carpeta_salida).mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    ws_cli = wb[hoja_clientes]
    ws_ind = wb[hoja_individual]

    # BUSCA LA CARPETA Y/O LA CREA
    carpeta_salida = Path(r"C:\info_juriscoop\Formato 341")
    carpeta_salida.mkdir(parents=True, exist_ok=True)
    
    # CREA ARCHIVO DE SALIDA
    nombre_txt = f"FTO_341_{trimestre_fin.strftime('%d%m%Y')}.txt"
    ruta_txt = carpeta_salida / nombre_txt
    txtFile = open(ruta_txt, "w", encoding="utf-8")

    
    #Encabezado
    Fila = 1
    
    linea1 = f"{Fila:08d}104{121:06d}{trimestre_fin.strftime('%d%m%Y')}{{UltimaFila}}CFCFINJURIS0300"
    txtFile.write(linea1 + "\n")
    Fila += 1
    print("Generando archivo plano 341, por favor espera...")

    #Registro 2
    print("Generando registros tipo 2...")
    for a in range(2, ws_cli.max_row + 1):
        TI = ws_cli[f"A{a}"].value or ""
        Id = ws_cli[f"B{a}"].value or ""
        CodVer = ws_cli[f"C{a}"].value or ""
        Nombre = ws_cli[f"D{a}"].value or ""
        if not TI or not Id or not CodVer or not Nombre:
            continue

        Nombres = str(Nombre)[:50] if len(str(Nombre)) > 50 else str(Nombre)
        id_str = str(Id)
        linea2 = (
            f"{Fila:08d}2{TI}"
            + "0" * max(0, 12 - len(id_str)) + id_str
            + str(CodVer)
            + Nombres
            + " " * (50 - len(Nombres))
        )
        txtFile.write(linea2 + "\n")
        Fila += 1

    # Registro 3
    txtFile.write(f"{Fila:08d}3" + "0" * 34 + "\n")
    Fila += 1

    # Registro 4 y 5
    print("Generando registros tipo 4 y 5")

    # Leer filas individuales
    filas_individual = []
    for r in range(fila_inicio, ws_ind.max_row + 1):
        UC = ws_ind[f"B{r}"].value
        if UC is None or str(UC).strip() == "":
            continue
        SubCta = ws_ind[f"C{r}"].value
        TI2 = ws_ind[f"D{r}"].value
        Id2 = ws_ind[f"E{r}"].value
        CIIU = ws_ind[f"H{r}"].value
        priNatu = ws_ind[f"I{r}"].value
        minRes = ws_ind[f"J{r}"].value
        priAcree = ws_ind[f"K{r}"].value
        priAcu = ws_ind[f"L{r}"].value
        valores = [ws_ind.cell(row=r, column=c).value for c in range(13, 53)]
        filas_individual.append((UC, SubCta, TI2, Id2, CIIU, priNatu, minRes, priAcree, priAcu, valores))
        
    conteos_por_id = Counter([row[3] for row in filas_individual if row[3]])

    idx = 0
    total_filas = len(filas_individual)

    while idx < total_filas:
        UC, SubCta, TI2, Id2, CIIU, priNatu, minRes, priAcree, priAcu, _ = filas_individual[idx]
        count_id = conteos_por_id.get(Id2, 1)
        fin_cliente = min(idx + count_id, total_filas)

        # Registro 4
        if TI2 not in (None, "", "0") and Id2 not in (None, "", "0"):
            TI2_s = str(TI2).strip()
            Id2_s = str(Id2).strip()
            CIIU_str = str(CIIU or "").strip()

            def safe_digit(x):
                s = str(x).strip() if x is not None else "0"
                return s if s.isdigit() else "0"

            priNatu_s, minRes_s, priAcree_s, priAcu_s = map(safe_digit, [priNatu, minRes, priAcree, priAcu])
            Ajuste = Id2_s[:9] if TI2_s.lower() == "nit" else Id2_s
            id_fill = "0" * max(0, 12 - len(Ajuste)) + Ajuste
            ciiu_fill = "0" * max(0, 4 - len(CIIU_str)) + CIIU_str

            linea4 = f"{Fila:08d}4{TI2_s}{id_fill}{ciiu_fill}{priNatu_s}{minRes_s}{priAcree_s}{priAcu_s}0"
            txtFile.write(linea4 + "\n")
            Fila += 1

        #Registro 5
        lineas_tipo5 = []              # lista de tuplas: (Columna_logica, UC, SubCta, valor, tipo_str)
        Suma_por_col_logica = {}      # key = columna_logica value = acumulado float

        for k in range(idx, fin_cliente):
            UC_k, SubCta_k, _, _, _, _, _, _, _, valores_k = filas_individual[k]
            for offset, Base_raw in enumerate(valores_k, start=13):
                Columna_logica = ws_ind.cell(row=1, column=offset).value
                if Columna_logica is None:
                    continue
                # convertir si es posible
                try:
                    valor_numerico = float(Base_raw) if Base_raw not in (None, "") else 0.0
                except (TypeError, ValueError):
                    valor_numerico = 0.0

                # incluir si tiene valor o si la columna está forzada (28, 29 o 40)
                if (valor_numerico != 0.0) or (Columna_logica in (28, 29, 40)):
                    # acumular para UC=13 (usamos la columna lógica como key)
                    Suma_por_col_logica.setdefault(Columna_logica, 0.0)
                    Suma_por_col_logica[Columna_logica] += valor_numerico

                    # decidir tipo (decimal o entero) según tu mapeo original
                    if (3 <= Columna_logica <= 14) or (Columna_logica in (17, 19, 21, 23, 28, 31, 32, 33, 34, 35, 36, 37, 38, 39, 40)):
                        tipo = "float"
                    else:
                        tipo = "int"

                    # guardamos el valor por fila (no numerado aún)
                    lineas_tipo5.append((int(Columna_logica), int(UC_k), int(SubCta_k), valor_numerico, tipo))

        # Generamos UC 13
        for col_logica, suma_val in Suma_por_col_logica.items():
            # Excluir columnas 25 a 28 del cálculo
            if col_logica in (25, 26, 27, 28):
                continue
        
            # Solo generar si hay valor o si es columna 29
            if (suma_val != 0.0) or (col_logica == 29):
                # Determinar si va con decimales o enteros
                if (3 <= col_logica <= 14) or col_logica in (17, 19, 21, 23, 32, 33, 34, 35, 36, 37, 38):
                    tipo = "float"
                elif (1 <= col_logica <= 2) or (col_logica == 29):
                    tipo = "int"
                else:
                    continue 
        
                lineas_tipo5.append((int(col_logica), 13, 5, suma_val, tipo))
                
        # Organizamos para hacer el Sort
        lineas_tipo5.sort(key=lambda t: (t[0], t[1], t[2]))

        for col_logica, uc_val, subcta_val, valor, tipo in lineas_tipo5:
            if tipo == "float":
                texto = f"{Fila:08d}5341{col_logica:02d}{uc_val:02d}{subcta_val:03d}+{abs(valor):017.2f}"
            else:
                texto = f"{Fila:08d}5341{col_logica:02d}{uc_val:02d}{subcta_val:03d}+{int(abs(valor)):017d}"
            txtFile.write(texto + "\n")
            Fila += 1

        # Avanza Cliente
        idx = fin_cliente

    #Registro 6
    txtFile.write(f"{Fila:08d}6\n")
    UltimaFila = Fila
    txtFile.close()

    # Actualiza encabezado con consecutivo
    with open(ruta_txt, "r", encoding="utf-8") as f:
        contenido = f.read().splitlines()

    Encabezado = linea1.replace("{UltimaFila}", f"{UltimaFila:08d}")
    contenido[0] = Encabezado

    with open(ruta_txt, "w", encoding="utf-8") as f:
        f.write("\n".join(contenido))

    print(f"\nArchivo plano 341 generado correctamente")
    print(f"Total de líneas: {UltimaFila}")

if __name__ == "__main__":
    Generar_Plano_341()
